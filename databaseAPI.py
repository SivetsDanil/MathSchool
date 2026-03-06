import os
import json
from typing import List, Dict
from openpyxl import load_workbook

# Путь к Excel-файлу
EXCEL_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'tutors.xlsx')

def get_tutors_by_subject(subject: str) -> List[Dict]:
    """Читает преподавателей из Excel-файла по предмету"""
    if not os.path.exists(EXCEL_PATH):
        return []

    try:
        wb = load_workbook(EXCEL_PATH, read_only=True)
        ws = wb.active

        # Читаем заголовки (первая строка)
        headers = [cell.value for cell in ws[1]]

        tutors = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            # Преобразуем строку в словарь
            tutor = dict(zip(headers, row))

            # Парсим tags: если это строка вида '["tag1", "tag2"]' → json.loads
            # Если заказчик будет писать через запятую → замени на: [t.strip() for t in tutor['tags'].split(',')]
            try:
                if isinstance(tutor.get('tags'), str):
                    tutor['tags'] = json.loads(tutor['tags'])
                else:
                    tutor['tags'] = []
            except (json.JSONDecodeError, TypeError):
                tutor['tags'] = []

            # Фильтруем по предмету
            if subject == '' or tutor.get('subject') == subject:
                tutors.append(tutor)

        wb.close()
        return tutors

    except Exception as e:
        print(f"Ошибка при чтении Excel: {e}")
        return []